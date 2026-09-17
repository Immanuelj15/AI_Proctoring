import io
import pytest
import docx
from unittest.mock import patch, AsyncMock
from app.services.pdf_extractor import (
    extract_text_from_docx,
    extract_text_from_file,
    generate_questions_with_distribution,
    parse_questions_heuristic,
)


def generate_mock_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("Cloud Computing & Distributed Systems Final Exam", level=1)
    doc.add_paragraph("1. Which cloud service model provides virtualized hardware resources over the internet? [2 Marks]")
    doc.add_paragraph("(A) SaaS  (B) PaaS  (C) IaaS  (D) FaaS")
    doc.add_paragraph("Answer: C")
    doc.add_paragraph("2. What is the difference between synchronous and asynchronous replication in distributed storage? [5 Marks]")
    doc.add_paragraph("3. Explain the CAP Theorem and discuss why partition tolerance is unavoidable in distributed databases. [10 Marks]")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_text_extraction():
    docx_bytes = generate_mock_docx_bytes()
    assert len(docx_bytes) > 0

    text = extract_text_from_docx(docx_bytes)
    assert "Cloud Computing" in text
    assert "CAP Theorem" in text
    assert "IaaS" in text


def test_parametric_difficulty_distribution():
    text = """
    1. Define latency in computer networking. [2 Marks] (A) Delay (B) Bandwidth (C) Throughput (D) Jitter. Answer: A
    2. What is a socket? [2 Marks] (A) Endpoint (B) Router (C) Cable (D) Switch. Answer: A
    3. Explain the three-way handshake in TCP. [5 Marks]
    4. Contrast distance-vector and link-state routing algorithms. [5 Marks]
    5. Design a fault-tolerant consensus algorithm using Raft. [10 Marks]
    6. Analyze Byzantine Fault Tolerance in decentralized systems. [10 Marks]
    """

    # Request exactly: 2 Easy, 3 Medium, 1 Hard
    questions = generate_questions_with_distribution(
        raw_text=text,
        subject="Networking",
        easy_count=2,
        medium_count=3,
        hard_count=1,
        allowed_types=["MCQ", "SHORT_ANSWER", "LONG_ANSWER"]
    )

    assert len(questions) == 6
    easy_qs = [q for q in questions if q["difficulty"] == "EASY"]
    medium_qs = [q for q in questions if q["difficulty"] == "MEDIUM"]
    hard_qs = [q for q in questions if q["difficulty"] == "HARD"]

    assert len(easy_qs) == 2
    assert len(medium_qs) == 3
    assert len(hard_qs) == 1


def test_api_extract_docx_file(client):
    # Setup examiner
    ex_reg = client.post("/auth/register", json={
        "name": "Prof. Distributed",
        "email": "distributed_prof@example.com",
        "password": "Password123!",
        "role": "examiner"
    })
    assert ex_reg.status_code == 201

    login_res = client.post("/auth/login", json={
        "email": "distributed_prof@example.com",
        "password": "Password123!"
    })
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    docx_bytes = generate_mock_docx_bytes()

    res = client.post(
        "/questions/extract-file",
        files={"file": ("systems_exam.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={
            "subject": "Distributed Systems",
            "easy_count": "1",
            "medium_count": "1",
            "hard_count": "1"
        },
        headers=headers
    )
    assert res.status_code == 200
    data = res.json()
    assert data["extracted_count"] == 3
    questions = data["questions"]
    assert any(q["difficulty"] == "EASY" for q in questions)
    assert any(q["difficulty"] == "MEDIUM" for q in questions)
    assert any(q["difficulty"] == "HARD" for q in questions)


@pytest.mark.asyncio
async def test_api_extract_url(client):
    client.post("/auth/register", json={
        "name": "Prof. Quantum",
        "email": "quantum_prof@example.com",
        "password": "Password123!",
        "role": "examiner"
    })

    login_res = client.post("/auth/login", json={
        "email": "quantum_prof@example.com",
        "password": "Password123!"
    })
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    mock_html = """
    <html>
      <body>
        <main>
          <h1>Introduction to Quantum Computing</h1>
          <p>Quantum computing leverages the phenomena of superposition and quantum entanglement to perform computations exponentially faster than classical computers for specific problem classes.</p>
          <p>Qubits serve as the fundamental unit of quantum information, able to exist in continuous states on the Bloch sphere.</p>
          <p>Shor's algorithm efficiently factors large integers in polynomial time, posing a fundamental challenge to RSA public key cryptography.</p>
        </main>
      </body>
    </html>
    """

    with patch("httpx.AsyncClient.get") as mock_get:
        mock_response = AsyncMock()
        mock_response.text = mock_html
        mock_response.raise_for_status = lambda: None
        mock_get.return_value = mock_response

        res = client.post(
            "/questions/extract-url",
            json={
                "url": "https://example.org/quantum-computing-lecture",
                "subject": "Quantum Computing",
                "easy_count": 1,
                "medium_count": 2,
                "hard_count": 1,
                "question_types": ["MCQ", "SHORT_ANSWER"]
            },
            headers=headers
        )

        assert res.status_code == 200
        data = res.json()
        assert data["extracted_count"] == 4
        assert len([q for q in data["questions"] if q["difficulty"] == "MEDIUM"]) == 2
