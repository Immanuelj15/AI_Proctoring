import io
import json
import os
import re
from typing import List, Optional
import pypdf
import docx
from bs4 import BeautifulSoup
import httpx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts and sanitizes text from raw PDF bytes using pypdf.
    """
    if not file_bytes:
        return ""

    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    extracted_pages = []

    for idx, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            extracted_pages.append(page_text.strip())

    return "\n\n".join(extracted_pages)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from Word documents (.docx) using python-docx.
    """
    if not file_bytes:
        return ""

    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = []

    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                paragraphs.append(" | ".join(row_texts))

    return "\n\n".join(paragraphs)


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """
    Multi-format file extractor supporting .pdf, .docx, .txt, and .md.
    """
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif lower_name.endswith(".txt") or lower_name.endswith(".md"):
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="ignore")
    else:
        # Try PDF first, then fallback to text
        try:
            return extract_text_from_pdf(file_bytes)
        except Exception:
            return file_bytes.decode("utf-8", errors="ignore")


async def fetch_and_clean_web_content(url: str) -> str:
    """
    Fetches external web page content via HTTP and extracts clean body text,
    stripping navigation, footer, script, and styling boilerplate.
    """
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.5",
    }

    async with httpx.AsyncClient(timeout=15.0, follow_redirects=True, headers=headers) as client:
        response = await client.get(url)
        response.raise_for_status()
        html_content = response.text

    soup = BeautifulSoup(html_content, "html.parser")

    # Remove non-content elements
    for element in soup(["script", "style", "nav", "footer", "header", "noscript", "aside", "form", "svg", "button"]):
        element.decompose()

    # Priority to article, main, or content containers
    main_container = soup.find("article") or soup.find("main") or soup.find("div", {"id": "content"}) or soup.find("div", {"class": "mw-parser-output"}) or soup.body

    if not main_container:
        main_container = soup

    lines = []
    for elem in main_container.find_all(["h1", "h2", "h3", "h4", "p", "li"]):
        text = elem.get_text(separator=" ", strip=True)
        if len(text) > 20:
            lines.append(text)

    cleaned_text = "\n\n".join(lines)
    return cleaned_text[:20000]  # Limit to 20,000 characters for token safety


def parse_questions_with_llm(
    raw_text: str,
    default_subject: str = "General",
    easy_count: Optional[int] = None,
    medium_count: Optional[int] = None,
    hard_count: Optional[int] = None,
    allowed_types: Optional[List[str]] = None
) -> Optional[List[dict]]:
    """
    Extracts or generates structured questions from text using OpenAI GPT-4o,
    enforcing requested difficulty distributions and question types.
    """
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key or not api_key.startswith("sk-"):
        return None

    try:
        import openai
        client = openai.OpenAI(api_key=api_key)

        types_clause = f"Allowed question types: {', '.join(allowed_types)}" if allowed_types else "Allowed question types: MCQ, SHORT_ANSWER, LONG_ANSWER"

        distribution_clause = ""
        if easy_count is not None and medium_count is not None and hard_count is not None:
            total_req = easy_count + medium_count + hard_count
            distribution_clause = f"""
            CRITICAL DIFFICULTY & COUNT REQUIREMENT:
            Generate/Extract exactly:
            - {easy_count} questions with "difficulty": "EASY"
            - {medium_count} questions with "difficulty": "MEDIUM"
            - {hard_count} questions with "difficulty": "HARD"
            Total questions to return: {total_req}.
            """
        else:
            distribution_clause = "Extract all clear questions identified in the text with appropriate difficulty levels (EASY, MEDIUM, or HARD)."

        prompt = f"""
        You are an expert assessment creator and examiner.
        Analyze the following study material, article, or question paper text and produce structured examination questions.

        {distribution_clause}
        {types_clause}
        Subject domain: "{default_subject}"

        For each question provide:
        - "question_text": Clean, academic question stem.
        - "question_type": "MCQ", "SHORT_ANSWER", or "LONG_ANSWER".
        - "subject": "{default_subject}".
        - "difficulty": "EASY", "MEDIUM", or "HARD".
        - "marks": 1.0 or 2.0 for MCQs, 5.0 for Short Answer, 10.0 for Long Answer.
        - "negative_marks": 0.25 if MCQ, else 0.0.
        - "model_answer": Detailed model answer or explanation key.
        - "options": For MCQ, provide exactly 4 options with "option_text" and "is_correct" (boolean, exactly ONE true). For non-MCQs, empty list [].

        Return ONLY a JSON object with the root key "questions": [...]

        SOURCE MATERIAL:
        {raw_text[:14000]}
        """

        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"},
            temperature=0.2
        )

        content = response.choices[0].message.content
        data = json.loads(content)
        questions = data.get("questions", [])
        if isinstance(questions, list) and len(questions) > 0:
            return questions
    except Exception as e:
        print(f"[pdf_extractor] LLM extraction error: {e}")

    return None


def parse_questions_heuristic(
    raw_text: str,
    default_subject: str = "General",
    easy_count: Optional[int] = None,
    medium_count: Optional[int] = None,
    hard_count: Optional[int] = None,
    allowed_types: Optional[List[str]] = None
) -> List[dict]:
    """
    Deterministic rule-based parser and parametric question synthesizer for offline mode.
    Handles numbered question papers AND extracts questions from article paragraphs
    matching requested difficulty distributions.
    """
    questions = []
    if not raw_text or not raw_text.strip():
        return questions

    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")

    # 1. Check if the text already contains pre-formatted questions
    q_split_pattern = r"(?=(?:^|\n)\s*(?:Q(?:uestion)?\s*\d+[\.\:\)]|\d+[\.\)])\s+)"
    raw_blocks = re.split(q_split_pattern, text, flags=re.IGNORECASE)

    extracted_from_paper = []
    for block in raw_blocks:
        block = block.strip()
        if not block:
            continue

        start_match = re.match(r"^(?:Q(?:uestion)?\s*\d+[\.\:\)]|\d+[\.\)])\s*", block, flags=re.IGNORECASE)
        if not start_match:
            continue

        content = block[start_match.end():].strip()
        if not content:
            continue

        marks = 2.0
        neg_marks = 0.0
        mark_match = re.search(r"[\[\(](\d+(?:\.\d+)?)\s*(?:marks?|pts?|points?)[\]\)]", content, re.IGNORECASE)
        if mark_match:
            try:
                marks = float(mark_match.group(1))
            except ValueError:
                pass
            content = content[:mark_match.start()] + content[mark_match.end():]

        detected_correct = None
        ans_match = re.search(r"(?:Answer|Ans|Correct(?:\s*Option)?)\s*[:\-]?\s*[\(\[]?([A-Da-d1-4])[\)\]]?", content, re.IGNORECASE)
        if ans_match:
            detected_correct = ans_match.group(1).upper()
            content = content[:ans_match.start()].strip()

        opt_pattern = r"(?:^|\n|\s{2,})\s*(?:\(?([A-Da-d])\)|([A-Da-d])[\.\)])\s+"
        opt_splits = list(re.finditer(opt_pattern, content))

        options = []
        q_type = "SHORT_ANSWER"
        q_text = content

        if len(opt_splits) >= 2:
            q_type = "MCQ"
            q_text = content[:opt_splits[0].start()].strip()
            for i, match in enumerate(opt_splits):
                letter = (match.group(1) or match.group(2)).upper()
                start_pos = match.end()
                end_pos = opt_splits[i + 1].start() if i + 1 < len(opt_splits) else len(content)
                opt_str = content[start_pos:end_pos].strip()

                is_corr = False
                if detected_correct and letter == detected_correct:
                    is_corr = True
                elif not detected_correct and i == 0:
                    is_corr = True

                options.append({"option_text": opt_str, "is_correct": is_corr})
        else:
            if marks >= 8.0 or len(q_text) > 180:
                q_type = "LONG_ANSWER"
            else:
                q_type = "SHORT_ANSWER"

        clean_q = re.sub(r"\s+", " ", q_text).strip()
        if len(clean_q) >= 5:
            extracted_from_paper.append({
                "question_text": clean_q,
                "question_type": q_type,
                "subject": default_subject,
                "difficulty": "MEDIUM",
                "marks": marks,
                "negative_marks": neg_marks,
                "model_answer": f"Answer key: Option {detected_correct}" if detected_correct else "Refer to question context",
                "options": options
            })

    # If the text is an article/study notes without explicit numbered questions, synthesize from paragraphs
    if len(extracted_from_paper) == 0:
        paragraphs = [p.strip() for p in text.split("\n\n") if len(p.strip()) > 40]
        for i, p in enumerate(paragraphs[:15]):
            sentences = [s.strip() for s in p.split(".") if len(s.strip()) > 15]
            if not sentences:
                continue

            first_sent = sentences[0]
            q_type = "MCQ" if (i % 2 == 0) else "SHORT_ANSWER"

            options = []
            if q_type == "MCQ":
                options = [
                    {"option_text": f"Accurate concept: {first_sent[:50]}", "is_correct": True},
                    {"option_text": "Inverted polarity principle", "is_correct": False},
                    {"option_text": "Non-applicable theoretical case", "is_correct": False},
                    {"option_text": "Legacy standard deprecated protocol", "is_correct": False},
                ]

            extracted_from_paper.append({
                "question_text": f"Based on {default_subject}: How does {first_sent[:70]}... apply?",
                "question_type": q_type,
                "subject": default_subject,
                "difficulty": "MEDIUM",
                "marks": 2.0 if q_type == "MCQ" else 5.0,
                "negative_marks": 0.5 if q_type == "MCQ" else 0.0,
                "model_answer": p[:200],
                "options": options
            })

    # If distribution counts (easy, medium, hard) are requested, partition and tag questions accordingly
    if easy_count is not None and medium_count is not None and hard_count is not None:
        total_needed = easy_count + medium_count + hard_count
        # Multiply/cycle available pool if fewer than needed
        pool = extracted_from_paper[:]
        while len(pool) < total_needed and len(pool) > 0:
            pool.extend([dict(item) for item in extracted_from_paper])

        final_list = []
        # Assign EASY
        for i in range(min(easy_count, len(pool))):
            item = dict(pool.pop(0))
            item["difficulty"] = "EASY"
            if allowed_types and item["question_type"] not in allowed_types:
                item["question_type"] = allowed_types[0]
            final_list.append(item)

        # Assign MEDIUM
        for i in range(min(medium_count, len(pool))):
            item = dict(pool.pop(0))
            item["difficulty"] = "MEDIUM"
            if allowed_types and item["question_type"] not in allowed_types:
                item["question_type"] = allowed_types[0]
            final_list.append(item)

        # Assign HARD
        for i in range(min(hard_count, len(pool))):
            item = dict(pool.pop(0))
            item["difficulty"] = "HARD"
            if item["question_type"] == "SHORT_ANSWER":
                item["question_type"] = "LONG_ANSWER"
                item["marks"] = max(item["marks"], 10.0)
            if allowed_types and item["question_type"] not in allowed_types:
                item["question_type"] = allowed_types[0]
            final_list.append(item)

        return final_list

    return extracted_from_paper


def generate_questions_with_distribution(
    raw_text: str,
    subject: str = "General",
    easy_count: Optional[int] = None,
    medium_count: Optional[int] = None,
    hard_count: Optional[int] = None,
    allowed_types: Optional[List[str]] = None
) -> List[dict]:
    """
    Unified question extractor and generator supporting difficulty distribution targets.
    Uses GPT-4o if available, else falls back to deterministic heuristic parsing.
    """
    if not raw_text or not raw_text.strip():
        return []

    # 1. Try GPT-4o with distribution constraints
    llm_res = parse_questions_with_llm(
        raw_text=raw_text,
        default_subject=subject,
        easy_count=easy_count,
        medium_count=medium_count,
        hard_count=hard_count,
        allowed_types=allowed_types
    )
    if llm_res and len(llm_res) > 0:
        return llm_res

    # 2. Fallback to heuristic parser
    return parse_questions_heuristic(
        raw_text=raw_text,
        default_subject=subject,
        easy_count=easy_count,
        medium_count=medium_count,
        hard_count=hard_count,
        allowed_types=allowed_types
    )


def parse_questions_from_pdf(
    file_bytes: bytes,
    default_subject: str = "General",
    easy_count: Optional[int] = None,
    medium_count: Optional[int] = None,
    hard_count: Optional[int] = None,
    allowed_types: Optional[List[str]] = None
) -> List[dict]:
    """
    Backward-compatible wrapper for PDF extraction with optional difficulty distribution.
    """
    text = extract_text_from_pdf(file_bytes)
    if not text.strip():
        return []

    return generate_questions_with_distribution(
        raw_text=text,
        subject=default_subject,
        easy_count=easy_count,
        medium_count=medium_count,
        hard_count=hard_count,
        allowed_types=allowed_types
    )
